#!/usr/bin/env python3
"""한국 기업 표준 견적서 엑셀 생성 스크립트.

Usage:
    python generate_quote_xlsx.py <specs.json> <route_eval.json> <quote_calc.json> <output_path.xlsx>

openpyxl 필요: pip install openpyxl
"""
import json
import sys
from datetime import date, timedelta

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side, numbers
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


# ── style constants ──────────────────────────────────────────────
THIN = Side(style="thin")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TITLE_FONT = Font(name="맑은 고딕", size=22, bold=True)
SECTION_FONT = Font(name="맑은 고딕", size=11, bold=True)
NORMAL_FONT = Font(name="맑은 고딕", size=10)
MONEY_FMT = '#,##0.00'
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

COL_WIDTHS = [6, 28, 24, 10, 8, 14, 18, 22]


def _apply_border(ws, row_start, row_end, col_start, col_end):
    for r in range(row_start, row_end + 1):
        for c in range(col_start, col_end + 1):
            ws.cell(r, c).border = BORDER_ALL


def _set_col_widths(ws):
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(specs: dict, route_eval: list, quotes: list) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "견적서"
    _set_col_widths(ws)

    rfq = specs.get("rfq_number", "N/A")
    today = date.today()
    valid_until = today + timedelta(days=30)

    # ── 1. Title ─────────────────────────────────────────────────
    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    title_cell = ws.cell(row, 1, "견  적  서")
    title_cell.font = TITLE_FONT
    title_cell.alignment = CENTER

    # ── 2. Header info ───────────────────────────────────────────
    row = 3
    headers_left = [
        ("견적번호", f"QT-{rfq}"),
        ("견적일자", today.strftime("%Y-%m-%d")),
        ("유효기간", valid_until.strftime("%Y-%m-%d")),
    ]
    headers_right = [
        ("수신", f"{specs.get('customer_company', '')}  {specs.get('customer_name', '')} 귀하"),
        ("RFQ No.", rfq),
        ("도면번호", specs.get("drawing_no", "")),
    ]
    for i, (label, value) in enumerate(headers_left):
        r = row + i
        ws.cell(r, 1, label).font = SECTION_FONT
        ws.cell(r, 1).alignment = RIGHT
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(r, 2, value).font = NORMAL_FONT
    for i, (label, value) in enumerate(headers_right):
        r = row + i
        ws.cell(r, 5, label).font = SECTION_FONT
        ws.cell(r, 5).alignment = RIGHT
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        ws.cell(r, 6, value).font = NORMAL_FONT

    # ── 3. 공급자 정보 ───────────────────────────────────────────
    row = 7
    ws.cell(row, 1, "공급자").font = SECTION_FONT
    supplier_info = [
        ("상호", "Dongwoo Dongam Technology (Wuxi)"),
        ("대표자", "(대표자명)"),
        ("연락처", "(연락처)"),
    ]
    for i, (label, value) in enumerate(supplier_info):
        r = row + i
        ws.cell(r, 2, label).font = NORMAL_FONT
        ws.cell(r, 2).alignment = RIGHT
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        ws.cell(r, 3, value).font = NORMAL_FONT

    # ── 4. 합계 금액 ────────────────────────────────────────────
    recommended = next(
        (q for q in quotes if q.get("status") == "recommended"),
        quotes[0] if quotes else None,
    )
    row = 11
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    if recommended:
        total = recommended["total_quote_cny"]
        total_cell = ws.cell(
            row, 1,
            f"합계금액:  ¥ {total:,.2f} CNY  (부가세 별도)"
        )
    else:
        total_cell = ws.cell(row, 1, "합계금액:  — (유효 루트 없음)")
    total_cell.font = Font(name="맑은 고딕", size=14, bold=True)
    total_cell.alignment = CENTER
    total_cell.fill = PatternFill("solid", fgColor="FFF2CC")
    _apply_border(ws, row, row, 1, 8)

    # ── 5. 품목 테이블 ──────────────────────────────────────────
    row = 13
    col_headers = ["No.", "품명", "규격 (공정)", "수량", "단위", "단가(CNY)", "공급가액(CNY)", "비고"]
    for c, h in enumerate(col_headers, 1):
        cell = ws.cell(row, c, h)
        cell.font = SECTION_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER_ALL

    lot_size = specs.get("lot_size", 0)
    for idx, q in enumerate(quotes, 1):
        row += 1
        is_rec = q.get("status") == "recommended"
        supply_amount = round(q["quote_per_pc_cny"] * lot_size, 2)
        remark = "추천" if is_rec else ("대안" if q.get("status") == "valid" else "참고")

        vals = [
            idx,
            specs.get("part_name", ""),
            q.get("display_name", q["route_name"]),
            lot_size,
            "EA",
            q["quote_per_pc_cny"],
            supply_amount,
            remark,
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            cell.font = NORMAL_FONT
            cell.border = BORDER_ALL
            if isinstance(v, float):
                cell.number_format = MONEY_FMT
                cell.alignment = RIGHT
            elif isinstance(v, int):
                cell.alignment = CENTER
            else:
                cell.alignment = LEFT

    # 합계 행
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    ws.cell(row, 1, "합    계").font = SECTION_FONT
    ws.cell(row, 1).alignment = CENTER
    if recommended:
        ws.cell(row, 6).font = NORMAL_FONT
        ws.cell(row, 7, recommended["total_quote_cny"]).font = SECTION_FONT
        ws.cell(row, 7).number_format = MONEY_FMT
        ws.cell(row, 7).alignment = RIGHT
    _apply_border(ws, row, row, 1, 8)

    # ── 6. 상세 조건 ────────────────────────────────────────────
    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1, "상세 조건").font = SECTION_FONT
    ws.cell(row, 1).fill = HEADER_FILL
    _apply_border(ws, row, row, 1, 8)

    details = [
        ("대상 부품", f"{specs.get('part_name', '')} / {specs.get('drawing_no', '')} / {specs.get('material', '')}"),
        ("유효 경화층 깊이 (ECD)", f"{specs.get('case_depth_low', '')}-{specs.get('case_depth_high', '')} mm @ HV550"),
        ("표면 경도", f"HRC {specs.get('surface_hardness_low', '')}-{specs.get('surface_hardness_high', '')}"),
        ("심부 경도", f"HRC {specs.get('core_hardness_low', '')}-{specs.get('core_hardness_high', '')}"),
        ("변형 한도", f"총 런아웃 ≤ {specs.get('distortion_limit_mm', '')} mm"),
        ("연삭 후처리", "없음 (열처리 후 연삭 수정 불가)" if not specs.get("grinding_after_ht") else "있음"),
        ("수량 / 월 소요", f"{specs.get('lot_size', '')} EA / 월 {specs.get('monthly_volume', 'N/A')} EA"),
    ]
    if recommended:
        details += [
            ("추천 공정", f"{recommended.get('display_name', recommended['route_name'])} ({recommended['equipment_id']})"),
            ("사이클 타임", f"{recommended['cycle_time_hr']} hr/batch"),
            ("배치 수", f"{recommended['batches']} batch(es)"),
        ]
    # lead time
    rec_route = next((r for r in route_eval if r.get("status") == "recommended"), None)
    if rec_route and rec_route.get("lead_time_assessment"):
        details.append(("리드타임", rec_route["lead_time_assessment"]))
    details.append(("검사 항목", "경도 검사 (표면/심부) + 외관 검사"))

    for label, value in details:
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row, 1, label).font = SECTION_FONT
        ws.cell(row, 1).alignment = RIGHT
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        ws.cell(row, 3, value).font = NORMAL_FONT
        ws.cell(row, 3).alignment = LEFT
        _apply_border(ws, row, row, 1, 8)

    # ── 7. 비고 / 특기사항 ──────────────────────────────────────
    row += 2
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1, "비고 / 특기사항").font = SECTION_FONT
    ws.cell(row, 1).fill = HEADER_FILL
    _apply_border(ws, row, row, 1, 8)

    missing = specs.get("missing_specs", [])
    notes = []
    if missing:
        notes.append("[ 고객 확인 필요 사항 ]")
        for m in missing:
            name = m.get("item_name", str(m))
            reason = m.get("reason", "")
            notes.append(f"  - {name}: {reason}" if reason else f"  - {name}")
    notes.append("")
    notes.append("[ 일반 조건 ]")
    notes.append("  - 본 견적은 러프 견적(rough quote)이며, 최종 단가는 시험 가공(trial) 후 확정됩니다.")
    notes.append("  - 납기: 첫 배치 PO 수령 후 조건 확인 필요")
    notes.append("  - 포장/방청 조건: 미정 (고객 사양 확인 필요)")
    notes.append("  - 금액 단위: CNY (부가세 별도)")

    for line in notes:
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.cell(row, 1, line).font = NORMAL_FONT
        ws.cell(row, 1).alignment = LEFT

    return wb


def main():
    if len(sys.argv) != 5:
        print(
            "Usage: python generate_quote_xlsx.py <specs.json> <route_eval.json> <quote_calc.json> <output.xlsx>",
            file=sys.stderr,
        )
        sys.exit(1)

    with open(sys.argv[1], encoding="utf-8") as f:
        specs = json.load(f)
    with open(sys.argv[2], encoding="utf-8") as f:
        route_eval = json.load(f)
    with open(sys.argv[3], encoding="utf-8") as f:
        quotes = json.load(f)

    wb = build_workbook(specs, route_eval, quotes)
    output_path = sys.argv[4]
    wb.save(output_path)
    print(f"견적서 생성 완료: {output_path}")


if __name__ == "__main__":
    main()
